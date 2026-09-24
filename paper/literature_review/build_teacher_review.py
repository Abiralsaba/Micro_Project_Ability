"""Create a standalone teacher-submission review from verified local source records.
Run with python-docx and reportlab installed. No network requests.
"""
from pathlib import Path
import json,re,html
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'teacher_submission';OUT.mkdir(exist_ok=True)
sources={s['id']:s for s in json.loads((ROOT/'SOURCE_ASSESSMENTS.json').read_text())}
raw=(ROOT/'TEACHER_REVIEW_SOURCE.md').read_text()
ids=list(dict.fromkeys(re.findall(r'\{\{([PS]\d{2})\}\}',raw)))
numbers={id:i+1 for i,id in enumerate(ids)}
def reference(id):
    s=sources[id]
    authors=s['authors'].split('; ')
    author=', '.join(authors) if len(authors)<=6 else authors[0]+' et al.'
    detail=[]
    if s.get('volume'):detail.append('vol. '+s['volume'])
    if s.get('issue'):detail.append('no. '+s['issue'])
    if s.get('pages'):detail.append('pp. '+s['pages'])
    if s.get('article'):detail.append('art. '+s['article'])
    detail.append(str(s['year']))
    special=' Workshop publication; review procedure not verified.' if id=='P26' else ''
    doi=(' doi: '+s['doi']+'.') if s.get('doi') else ''
    return f'[{numbers[id]}] {author}, “{s["title"]},” {s["venue"]}, '+', '.join(detail)+'.'+doi+' '+f'[Official source]({s["url"]}).'+special
refs='\n\n'.join(reference(id) for id in ids)
access=['| Reference | Source type | Evidence examined |','|---|---|---|']
for id in ids:
    s=sources[id]
    access.append(f'| [{numbers[id]}] | {s["kind"]} | {s["level"]} |')
text=re.sub(r'\{\{([PS]\d{2})\}\}',lambda m:f'[{numbers[m[1]]}]({sources[m[1]]["url"]})',raw)
text=text.replace('{{REFERENCES}}',refs).replace('{{ACCESS}}','\n'.join(access))
assert '{{' not in text
(OUT/'AMP_Literature_Review.md').write_text(text)

def blocks(md):
    lines=md.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            table=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',v) for v in row):table.append(row)
                i+=1
            yield ('table',table);continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#'));yield('h'+str(level),line[level:].strip());i+=1;continue
        para=[line];i+=1
        while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|')):
            para.append(lines[i].strip());i+=1
        yield('p',' '.join(para))

LINK=re.compile(r'\[([^\]]+)\]\(([^)]+)\)')
def plain(s):
    return re.sub(r'\*\*(.*?)\*\*',r'\1',LINK.sub(lambda m:('['+m[1]+']') if m[1].isdigit() else m[1]+' ('+m[2]+')',s))

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
sec.top_margin=sec.bottom_margin=Inches(.8)
sec.left_margin=sec.right_margin=Inches(.85)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(7)
for name in ['Title','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor.from_string('17365D')
    st.paragraph_format.keep_with_next=True
doc.styles['Heading 1'].font.size=Pt(14)
doc.styles['Heading 2'].font.size=Pt(12)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('AMP Literature Review  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='Reliable Cross-Modal Assistive Communication: Literature Review'
doc.core_properties.subject='Presentation progress, receiver pacing and interruption recovery'
doc.core_properties.author=''

def add_runs(p,s):
    # Numeric references retain their links in the PDF/Markdown; Word includes readable source URLs.
    s=LINK.sub(lambda m:('['+m[1]+']') if m[1].isdigit() else m[2],s)
    for i,part in enumerate(re.split(r'\*\*(.*?)\*\*',s)):
        run=p.add_run(part);run.bold=bool(i%2)

fonts=Path('/System/Library/Fonts/Supplemental')
for name,file in [('ReviewSerif','Times New Roman.ttf'),('ReviewSerifBold','Times New Roman Bold.ttf'),('ReviewSerifItalic','Times New Roman Italic.ttf'),('ReviewSerifBoldItalic','Times New Roman Bold Italic.ttf')]:
    pdfmetrics.registerFont(TTFont(name,str(fonts/file)))
pdfmetrics.registerFontFamily('ReviewSerif',normal='ReviewSerif',bold='ReviewSerifBold',italic='ReviewSerifItalic',boldItalic='ReviewSerifBoldItalic')
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='BodyReview',fontName='ReviewSerif',fontSize=11,leading=14,spaceAfter=7,alignment=TA_JUSTIFY))
styles.add(ParagraphStyle(name='SmallReview',fontName='ReviewSerif',fontSize=9.5,leading=12,spaceAfter=6))
styles.add(ParagraphStyle(name='CellReview',fontName='ReviewSerif',fontSize=9,leading=11))
styles.add(ParagraphStyle(name='HeadingReview',fontName='ReviewSerifBold',fontSize=13.5,leading=17,spaceBefore=13,spaceAfter=7,keepWithNext=True,textColor=colors.HexColor('#17365D')))
styles.add(ParagraphStyle(name='SubReview',fontName='ReviewSerifBold',fontSize=11.5,leading=15,spaceBefore=9,spaceAfter=6,keepWithNext=True))
styles.add(ParagraphStyle(name='TitleReview',fontName='ReviewSerifBold',fontSize=19,leading=23,spaceAfter=15,textColor=colors.HexColor('#17365D')))
def pdf_markup(s):
    tokens={}
    def hold(m):
        key=f'LINKTOKEN{len(tokens)}END'
        label='['+m[1]+']' if m[1].isdigit() else m[1]
        tokens[key]=f'<link href="{html.escape(m[2],quote=True)}" color="#17365D">{html.escape(label)}</link>'
        return key
    s=LINK.sub(hold,s);s=html.escape(s)
    s=re.sub(r'\*\*(.*?)\*\*',r'<b>\1</b>',s)
    for k,v in tokens.items():s=s.replace(k,v)
    return s

story=[];in_refs=False
for kind,value in blocks(text):
    if kind.startswith('h'):
        if value in ['Abstract','References']:
            doc.add_page_break();story.append(PageBreak())
        in_refs=value=='References' or (in_refs and value!='Appendix A. Evidence access')
        if value=='Appendix A. Evidence access':in_refs=False
        level=int(kind[1:]);p=doc.add_heading(value,0 if level==1 else level-1)
        story.append(Paragraph(pdf_markup(value),styles['TitleReview' if level==1 else 'HeadingReview' if level==2 else 'SubReview']))
    elif kind=='table':
        table=doc.add_table(rows=1,cols=len(value[0]));table.style='Light Shading Accent 1'
        for j,v in enumerate(value[0]):add_runs(table.rows[0].cells[j].paragraphs[0],v)
        repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in value[1:]:
            cells=table.add_row().cells
            for j,v in enumerate(row):
                add_runs(cells[j].paragraphs[0],v)
                for run in cells[j].paragraphs[0].runs:run.font.size=Pt(9)
        widths=[100,142,231] if len(value)==10 else [65,190,218]
        pt=Table([[Paragraph(pdf_markup(c),styles['CellReview']) for c in row] for row in value],colWidths=widths,repeatRows=1,hAlign='LEFT')
        pt.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E8EDF3')),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#AAB6C3')),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
        story.extend([pt,Spacer(1,10)])
    else:
        p=doc.add_paragraph();add_runs(p,value)
        if in_refs:
            for run in p.runs:run.font.size=Pt(9)
        elif not value.startswith('**'):p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        story.append(Paragraph(pdf_markup(value),styles['SmallReview' if in_refs else 'BodyReview']))

doc.save(OUT/'AMP_Literature_Review.docx')
def page_footer(canvas,document):
    canvas.saveState();canvas.setFont('ReviewSerif',9);canvas.setFillColor(colors.HexColor('#555555'))
    canvas.drawCentredString(A4[0]/2,27,f'AMP Literature Review  |  {document.page}');canvas.restoreState()
pdf=SimpleDocTemplate(str(OUT/'AMP_Literature_Review.pdf'),pagesize=A4,rightMargin=61,leftMargin=61,topMargin=48,bottomMargin=48,title='Reliable Cross-Modal Assistive Communication: Literature Review',author='')
pdf.build(story,onFirstPage=page_footer,onLaterPages=page_footer)
manifest={'date':'2026-09-16','literature_cutoff':'2026-09-15','reference_count':len(ids),'reference_mapping':numbers,'body_word_count':len(text.split('## References')[0].split()),'note':'No experimental results claimed. Student metadata fields left blank for completion.'}
(OUT/'BUILD_INFO.json').write_text(json.dumps(manifest,indent=2))
print(json.dumps(manifest,indent=2))
