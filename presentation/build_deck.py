"""Build an offline animated deck and an editable PowerPoint from one scene model.
Run with Python 3 + python-pptx. No project firmware is imported or changed.
"""
from pathlib import Path
import html, json, csv
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parent
NAVY='101E2B'; PANEL='1B2D3C'; INK='152C3B'; CREAM='F5F2EA'; WHITE='FFFFFF'
MUTED='B6C6CF'; TEAL='77E1CB'; ORANGE='FFB17B'; PURPLE='C9B9FF'; GRAY='526775'
TEAM='Team: [names / IDs]  ·  [department / institution]'
slides=[]

class Slide:
    def __init__(self, title, chapter, light=False, notes='', appendix=False):
        self.title=title; self.chapter=chapter; self.light=light; self.notes=notes
        self.bg=CREAM if light else NAVY; self.fg=INK if light else WHITE
        self.sub=GRAY if light else MUTED; self.elements=[]; self.appendix=appendix
        slides.append(self)
        self.text(76,43,550,28,'ABILITY / PROJECT SHOW',18,color=self.sub,bold=True)
        self.text(1040,43,480,28,chapter.upper(),17,color=self.sub,align='right')
        self.line(76,814,1524,814, 'D7DDD9' if light else '344652',1)
        self.text(76,835,1240,24,'PROJECT ABILITY  /  Different ways to communicate. One connection.',16,color=self.sub)
        self.text(1430,835,90,24,f'{len(slides):02d}',18,color=self.sub,align='right')
    def text(self,x,y,w,h,t,size=28,color=None,bold=False,align='left',group=0,link=None):
        self.elements.append(dict(kind='text',x=x,y=y,w=w,h=h,text=t,size=size,color=color or self.fg,bold=bold,align=align,group=group,link=link))
    def box(self,x,y,w,h,fill=None,stroke=None,r=20,group=0,sw=2):
        self.elements.append(dict(kind='box',x=x,y=y,w=w,h=h,fill=fill,stroke=stroke,r=r,group=group,sw=sw))
    def circle(self,x,y,w,h=None,fill=None,stroke=None,sw=3,group=0):
        self.elements.append(dict(kind='circle',x=x,y=y,w=w,h=h or w,fill=fill,stroke=stroke,sw=sw,group=group))
    def line(self,x,y,x2,y2,color=None,sw=3,group=0):
        self.elements.append(dict(kind='line',x=x,y=y,x2=x2,y2=y2,color=color or self.sub,sw=sw,group=group))
    def head(self,kicker,title,subtitle=None):
        self.text(78,118,1430,28,kicker,20,color=GRAY if self.light else TEAL,bold=True)
        self.text(76,165,1450,146,title,64,bold=True)
        if subtitle:self.text(80,303,1420,62,subtitle,27,color=self.sub)
    def chip(self,x,y,label,color=TEAL,w=220,group=0):
        self.box(x,y,w,40,fill=None,stroke=color,r=19,group=group)
        self.text(x+12,y+9,w-24,26,label,16,color=color,bold=True,align='center',group=group)
    def card(self,x,y,w,h,n,title,body,color=TEAL,group=1):
        self.box(x,y,w,h,fill='FFFFFF' if self.light else PANEL,group=group)
        self.text(x+26,y+24,w-52,45,n,30,color=color,bold=True,group=group)
        self.text(x+26,y+90,w-52,90,title,34,bold=True,group=group)
        self.text(x+26,y+195,w-52,h-215,body,25,color=self.sub,group=group)

def arrow(s,x,y,x2,y2,color=TEAL,group=1):
    s.line(x,y,x2,y2,color,3,group)
    if abs(x2-x)>=abs(y2-y):
        d=1 if x2>x else -1
        s.line(x2,y2,x2-d*12,y2-8,color,3,group);s.line(x2,y2,x2-d*12,y2+8,color,3,group)
    else:
        d=1 if y2>y else -1
        s.line(x2,y2,x2-8,y2-d*12,color,3,group);s.line(x2,y2,x2+8,y2-d*12,color,3,group)
    s.elements.append(dict(kind='flow',x=x,y=y,x2=x2,y2=y2,color=color,group=group))

def dots(s,x,y,pattern=63,scale=1,color=TEAL,group=1):
    for c in range(2):
        for r in range(3):
            active=pattern & (1<<(c*3+r))
            s.circle(x+c*34*scale,y+r*34*scale,20*scale,fill=color if active else ('D2DAD6' if s.light else '38505E'),group=group)

def person(s,x,y,color=TEAL,scale=1,chair=False,group=1):
    s.circle(x+36*scale,y,66*scale,fill=ORANGE,group=group)
    s.box(x+20*scale,y+80*scale,95*scale,130*scale,fill=color,r=36,group=group)
    if chair:
        s.line(x+54*scale,y+200*scale,x+128*scale,y+216*scale,s.fg,15*scale,group)
        s.line(x+128*scale,y+216*scale,x+145*scale,y+269*scale,s.fg,15*scale,group)
        s.line(x+145*scale,y+269*scale,x+167*scale,y+269*scale,s.fg,12*scale,group)
    else:
        s.line(x+36*scale,y+204*scale,x+18*scale,y+296*scale,s.fg,13*scale,group)
        s.line(x+99*scale,y+204*scale,x+121*scale,y+296*scale,s.fg,13*scale,group)
    s.line(x+28*scale,y+110*scale,x-16*scale,y+160*scale,ORANGE,16*scale,group)
    s.line(x+107*scale,y+110*scale,x+152*scale,y+73*scale,ORANGE,16*scale,group)
    if chair:
        s.circle(x-22*scale,y+174*scale,132*scale,stroke=PURPLE,sw=9*scale,group=group)
        s.circle(x+131*scale,y+268*scale,34*scale,stroke=PURPLE,sw=6*scale,group=group)
        s.line(x+44*scale,y+232*scale,x+140*scale,y+232*scale,PURPLE,7*scale,group)
        s.line(x+140*scale,y+232*scale,x+147*scale,y+283*scale,PURPLE,7*scale,group)

def glove(s,x,y,scale=1,color=TEAL):
    s.box(x+35*scale,y+94*scale,132*scale,144*scale,fill=color,r=36,group=1)
    for i,off in enumerate([25,0,12,39]):
        s.box(x+(36+i*34)*scale,y+off*scale,29*scale,(125-off)*scale,fill=color,r=13,group=1)
        s.line(x+(50+i*34)*scale,y+(off+20)*scale,x+(50+i*34)*scale,y+146*scale,INK,4*scale,1)
    s.line(x+48*scale,y+155*scale,x+4*scale,y+113*scale,color,31*scale,1)
    s.box(x+64*scale,y+197*scale,73*scale,58*scale,fill=INK,r=9,group=1)
    s.text(x+65*scale,y+213*scale,72*scale,30*scale,'IMU',17*scale,color=WHITE,bold=True,align='center',group=1)

# 1 — emotional opening, original vector illustration
s=Slide('Everyone deserves a way to say hello.','01 / The reason',notes='[20 seconds] Imagine having something important to say, but the person beside you cannot receive it in the way you express it. A simple hello should not depend on sight, hearing, or steady hands. We are [team names], and this is Project Ability: different ways to communicate, one connection. Replace the editable team line before recording.')
s.text(78,139,880,52,'PROJECT ABILITY',31,color=TEAL,bold=True)
s.text(76,220,940,272,'A thought. A feeling.\nA simple “hello.”',83,bold=True)
s.text(80,540,770,116,'Everyone deserves a way\nto share it.',43,color=MUTED)
s.text(80,743,1380,34,TEAM,21,color=MUTED)
s.circle(1090,159,365,fill=PANEL)
person(s,1060,275,TEAL,0.9);person(s,1310,362,PURPLE,0.85,True)
s.box(1040,147,281,85,fill=TEAL,r=25,group=2)
s.text(1058,166,240,45,'hello.',40,color=INK,bold=True,align='center',group=2)
s.box(1230,650,245,90,fill=PANEL,stroke='38505E',group=3)
dots(s,1252,662,0x13,.6);dots(s,1310,662,0x0A,.6)
s.text(1370,675,90,38,'hi',31,color=TEAL,group=3)

# 2
s=Slide('Three interfaces. One shared conversation.','02 / The idea',light=True,notes='[25 seconds] Our idea connects personal interfaces through one shared hub. The glove turns selected hand gestures into messages. The Braille unit lets a user enter characters and feel received text. Our proposed third module adds a gaze keyboard and brain or eye-based wheelchair commands. The common language inside the system is text. The avatar, full return path and mobility extension still need integration; this is the ecosystem direction, not a claim that every route is deployed.')
s.head('THE PROBLEM → OUR IDEA','Three interfaces.\nOne shared conversation.')
s.card(78,356,462,347,'01 / SIGN','Glove + avatar','Gesture input for Deaf users.\nSigning-avatar output planned.','147C6F',1)
s.card(569,356,462,347,'02 / TOUCH','Braille + buttons','Tactile reading and character\nentry for blind users.','88633F',2)
s.card(1060,356,462,347,'03 / INTENT','Brain + eyes','Proposed mobility and messaging\nfor users with motor limitations.','7960AA',3)
s.text(80,742,1400,36,'The contribution: connect affordable interfaces through a common message format.',29,color=INK,bold=True)

# 3
s=Slide('Our own AI/ML processing hub.','03 / Architecture',notes='[30 seconds] Our platform vision is an AI/ML-based processing hub built on Raspberry Pi 5, using our own task-specific recognition model and an accessible chat interface. It brings model inference, message routing, receiver selection and format conversion into one place. The present repository implements routing and Braille encoding; custom model inference, chat history, avatar and speech remain integration milestones. Today the glove classifier runs on its ESP32. This diagram presents the intended architecture, not a claim that all these features are deployed.')
s.head('OUR PLATFORM VISION','Our own AI/ML processing hub.','Raspberry Pi 5 + custom recognition model + an accessible chat interface')
for y,title,body,col in [(408,'GLOVE','5 flex sensors + BNO055',TEAL),(538,'BRAILLE KEYS','6 dot buttons + action',ORANGE),(668,'GAZE / EEG','Proposed input interfaces',PURPLE)]:
    s.box(80,y,370,104,fill=PANEL,group=1)
    s.text(103,y+16,328,32,title,24,color=col,bold=True,group=1)
    s.text(103,y+55,328,27,body,20,color=MUTED,group=1)
    arrow(s,460,y+52,581,561,col,2)
s.box(600,402,408,326,fill=TEAL,r=30,group=2)
s.text(627,426,353,35,'AI/ML PROCESSING HUB',21,color=INK,bold=True,align='center',group=2)
s.text(623,478,363,64,'Raspberry Pi 5',42,color=INK,bold=True,align='center',group=2)
s.text(627,557,353,154,'Our own AI/ML model*\nChat + conversation history*\nMessage routing\nAccessible format conversion',24,color=INK,align='center',group=2)
for y,title,col in [(410,'Physical Braille',ORANGE),(535,'Chat interface*',TEAL),(660,'Avatar / speech*',PURPLE)]:
    arrow(s,1015,561,1137,y+45,col,3)
    s.box(1150,y,370,92,fill=PANEL,group=3)
    s.text(1172,y+26,329,45,title,29,color=col,bold=True,group=3)
s.text(602,749,920,45,'*Proposed integration. Routing + Braille encoding exist in the current hub code.',17,color=MUTED)

# 4
s=Slide('A hand movement becomes a message.','04 / Glove module',notes='[25 seconds] Five flex sensors measure finger bend, and a BNO055 measures orientation. The ESP32-S3 calibrates the values, applies gesture rules and builds text. A send gesture publishes the message to the Raspberry Pi. This is a limited gesture and fingerspelling prototype, not full sign-language translation. Our next recognition step is a custom-trained sensor-and-vision model. A signing avatar is also part of the planned receive interface, but its renderer and clips are absent from this checkout.')
s.head('FEATURE 01 / GESTURE INPUT','A hand movement\nbecomes a message.')
s.box(80,359,451,409,fill=PANEL)
glove(s,189,402,1.12)
s.text(107,705,397,40,'5 flex sensors + BNO055 IMU',23,color=TEAL,bold=True,align='center')
for i,(a,b) in enumerate([('SENSE','Finger bend + palm orientation'),('INTERPRET','ESP32-S3 calibration + gesture rules'),('SEND','Message → MQTT → Raspberry Pi')]):
    y=364+i*101
    s.text(598,y,208,40,a,25,color=TEAL,bold=True,group=i+1)
    s.text(809,y,688,69,b,29,color=WHITE,group=i+1)
s.box(589,683,931,88,fill=PANEL,stroke=PURPLE,group=4)
s.text(616,699,884,64,'NEXT: our own trained fusion model + a signing-avatar receiver',27,color=PURPLE,bold=True,group=4)

# 5
s=Slide('Read with touch. Reply with buttons.','05 / Braille module',light=True,notes='[30 seconds] The latest Braille firmware has four six-dot cells with eight MG90S servos: two servos per cell. Cam mechanisms convert servo position into dot patterns. Six dot buttons accumulate a character, and the action button commits, spaces or displays the buffer. Show the actual button sequence during the demo. In the current cloud firmware, triple-click sends the buffer to its own display, not back through MQTT. Sending that reply to the glove or avatar remains an integration step. The letters shown here are the correct six-dot patterns for HELP; physical correctness still needs verification.')
s.head('FEATURES 02 + 03 / TACTILE OUTPUT + BUTTON INPUT','Read with touch.\nReply with buttons.')
s.box(80,359,845,282,fill=WHITE)
for i,(ch,pat) in enumerate(zip('HELP',[0x13,0x11,0x07,0x0F])):
    x=112+i*201
    s.box(x,387,170,181,fill=CREAM,r=17,group=i+1)
    dots(s,x+51,411,pat,1.2,'147C6F',i+1)
    s.text(x,576,170,49,ch,32,color=INK,bold=True,align='center',group=i+1)
s.text(104,667,800,45,'4 cells × 2 servos = 8 MG90S actuators',32,color=INK,bold=True)
s.text(104,725,800,43,'Cam mechanism • physical dots • replaceable parts',24,color=GRAY)
s.box(961,359,559,411,fill=INK)
s.text(990,387,505,48,'6 DOT KEYS + 1 ACTION',27,color=TEAL,bold=True)
for i in range(6):
    s.circle(1005+i*77,472,52,fill=TEAL,group=2)
    s.text(1005+i*77,483,52,30,str(i+1),23,color=INK,bold=True,align='center',group=2)
s.text(990,566,505,135,'Tap dots → commit character\nDouble-click → space\nTriple-click → local display',27,color=WHITE,group=3)
s.text(990,718,502,32,'Hub reply / avatar link: next integration',19,color=ORANGE)

# 6
s=Slide('Small moments. Real independence.','06 / Real-life impact',notes='[20 seconds] We are targeting practical moments: a classroom partner sharing a short message, a family communicating across different access needs, and eventually a person with motor limitations composing a request with their eyes. These are intended use scenarios. The benefits we want to measure are successful conversations, fewer access barriers and easier interaction. People have different abilities and preferences, so the appropriate interface must be selected with each user.')
s.head('WHERE IT MATTERS','Small moments.\nReal independence.')
s.card(80,355,460,370,'CLASSROOM','“Your turn.”','A gesture becomes tactile text\nfor a blind classmate.',TEAL,1)
s.card(570,355,460,370,'HOME','“I’m here.”','Family members exchange\nmessages in accessible formats.',ORANGE,2)
s.card(1060,355,460,370,'CARE / PROPOSED','“Water, please.”','A gaze keyboard composes\na request without hand typing.',PURPLE,3)
s.text(81,759,1440,34,'Designed for Deaf, blind, speech-impaired and motor-impaired users, plus their communication partners.',22,color=MUTED)

# 7
s=Slide('Now, follow one message.','07 / Live demo handoff',light=True,notes='[10 seconds, then 2 minutes 45 seconds LIVE HARDWARE, NO SLIDES] Say: Now let us follow one message through the real hardware. End screen sharing or press B to black out the browser. Follow DEMO_SCRIPT.md: show boards/power, glove input, hub receipt and four tactile cells; then show Braille button entry and local output. Use a text injection only as an explicitly labeled transport test. Do not represent keyboard-to-avatar or wheelchair animation as working hardware. After the demonstration return directly to slide 08. This is a handoff, not a slide substituting for the demo.')
s.text(80,145,1250,39,'THE MOST IMPORTANT PART',21,color='147C6F',bold=True)
s.text(75,228,1430,201,'Now, follow\none message.',88,bold=True)
for i,(a,b) in enumerate([('01','Move a hand'),('02','Watch the Pi route it'),('03','Feel the Braille')]):
    x=82+i*485
    s.text(x,530,425,65,a,48,color='147C6F',bold=True,group=i+1)
    s.text(x,621,430,75,b,34,color=INK,bold=True,group=i+1)
s.text(81,751,1400,37,'LIVE HARDWARE DEMONSTRATION →',26,color=GRAY,bold=True)

# 8
s=Slide('Built for a local budget.','08 / Cost effectiveness',notes='[30 seconds] All amounts are in Bangladeshi taka. You supplied a glove estimate of 4,000 to 5,000 taka; we use 4,500 as the working midpoint. The other amounts are design-to-cost targets, not receipts: Braille 3,500, voice 1,500, gaze 3,500 and shared Pi hub 15,000. The core glove, Braille and hub target is 23,000. Adding voice and gaze brings the communication-system target to 28,000, or 30,800 with a ten-percent reserve. The lower hub target assumes a lower-memory or reused Pi 5 and an existing screen; model performance needs benchmarking. The optional wheelchair controller target of 12,000 excludes EEG acquisition and the powered chair. It is not the cost of a full brain-controlled wheelchair. No tax, shipping, labor or clinical validation cost is included.')
s.head('LOW-COST DESIGN / ALL VALUES IN BDT','Built for a local budget.')
items=[('GLOVE','৳4,500',TEAL,'Estimated range: ৳4,000–5,000'),('4-CELL BRAILLE','৳3,500',ORANGE,'Target: servos + keys + mechanics'),('VOICE MODULE','৳1,500',PURPLE,'Target: mic + speaker + ESP32'),('GAZE INTERFACE','৳3,500',TEAL,'Target: camera; screen + Pi shared'),('SHARED PI HUB','৳15,000',ORANGE,'Target: lower-memory / reused Pi 5'),('WHEELCHAIR ADD-ON','৳12,000*',PURPLE,'Control target; EEG + chair excluded')]
for i,(label,num,col,small) in enumerate(items):
    x=80+(i%3)*489;y=303+(i//3)*188
    s.box(x,y,460,170,fill=PANEL,group=i//3+1)
    s.text(x+24,y+17,417,30,label,21,color=col,bold=True,group=i//3+1)
    s.text(x+21,y+54,420,65,num,49,bold=True,group=i//3+1)
    s.text(x+24,y+129,417,32,small,20,color=MUTED,group=i//3+1)
s.text(82,700,835,57,'Communication target: ৳28,000',37,bold=True)
s.text(82,758,880,38,'Core: ৳23,000 · With voice/gaze + 10% reserve: ৳30,800',22,color=MUTED)
s.box(998,697,522,98,fill=TEAL,group=3)
s.text(1020,711,478,39,'Shared hub: ৳3,000 / user',30,color=INK,bold=True,group=3)
s.text(1020,754,478,28,'5 users · *Mobility budget is separate',20,color=INK,group=3)

# 9
s=Slide('Premium prices. A lower-cost direction.','09 / Market context',light=True,notes='[30 seconds] The price barrier is substantial. At the rounded presentation conversion of 123 taka per US dollar, Rokoko Smartgloves II cost about 2.45 lakh taka for a professional motion-capture pair. Our single gesture-glove estimate is 4,000 to 5,000 taka. HumanWare Brailliant BI 40X is about 4.62 lakh taka for a commercial 40-cell reader, while our four-cell prototype target is 3,500 taka. These are selected premium references, not the entire market or equivalent-performance comparisons. The product pages opened today show 1,995 and 3,759 US dollars respectively; an older search snippet showed a HumanWare sale price that the current page did not. Lower-cost commercial options also exist: Orbit Reader 20 is about 98,300 taka. Ability costs exclude a personal allocation of the shared hub, production support, tax and shipping. The affordability strategy is focused functions, shared processing and locally assembled parts.')
s.head('SELECTED PREMIUM REFERENCES / BDT','Premium prices.\nA lower-cost direction.')
for x,title,product,price,ours,scope,col in [(80,'HAND INPUT','Rokoko Smartgloves II','≈ ৳2.45 lakh','Ability: ৳4,000–5,000','Pro motion-capture pair vs our single\ngesture-glove prototype.','147C6F'),(820,'TACTILE OUTPUT','HumanWare Brailliant BI 40X','≈ ৳4.62 lakh','Ability target: ৳3,500','40-cell commercial reader vs our\n4-cell, six-dot prototype.','88633F')]:
    s.box(x,353,700,341,fill=WHITE,group=1)
    s.text(x+25,373,650,29,title,20,color=col,bold=True,group=1)
    s.text(x+25,416,650,38,product,29,color=INK,bold=True,group=1)
    s.text(x+22,465,653,73,price,51,color=GRAY,bold=True,group=1)
    s.box(x+24,548,652,66,fill=INK,r=13,group=2)
    s.text(x+40,563,620,44,ours,32,color=TEAL,bold=True,group=2)
    s.text(x+25,631,650,61,scope,23,color=GRAY,group=3)
s.text(82,715,1430,42,'Lower entry cost through focused functions, shared computing and local assembly.',28,color=INK,bold=True)
s.text(82,760,1420,49,'Price context, not equal performance. Module costs exclude hub share. USD 1 ≈ ৳123; tax/shipping excluded.\nOfficial listings checked 30 Sep 2026. Lower-price reference: Orbit Reader 20 ≈ ৳98,300 (20 cells).',17,color=GRAY)

# 10
s=Slide('The next module: mobility + a voice.','10 / Proposed extension',notes='[40 seconds] This is our proposed brain-and-eye wheelchair module, shown only as a concept. It targets people who need an alternative to hand controls, including some people with Parkinson’s after individual assessment. A multichannel EEG interface would classify a small set of trained intentions; a camera-based gaze interface offers another way to choose directions. In communication mode, the chair is stopped and a gaze keyboard sends text through the same Raspberry Pi to Braille, text, and eventually avatar or speech. A separate local controller must enforce an emergency stop, obstacle stop, command timeout and speed limit. It should never depend on cloud connectivity for stopping. Gaze calibration and real EEG classification remain research work. This is neither a medical claim nor demonstrated wheelchair hardware.')
s.head('PRESENTATION-ONLY CONCEPT','The next module:\nmobility + a voice.')
s.chip(1100,125,'PROPOSED / NOT BUILT',PURPLE,415)
s.box(80,357,450,407,fill=PANEL)
s.text(102,374,407,28,'PARKINSON’S / MOTOR ACCESS',19,color=PURPLE,bold=True)
person(s,200,401,TEAL,.95,True)
s.line(230,415,297,415,PURPLE,7,1)
s.circle(247,406,13,fill=PURPLE);s.circle(280,406,13,fill=PURPLE)
s.box(355,491,151,137,fill=NAVY,stroke=TEAL,r=12,group=2)
s.text(366,502,130,31,'WATER',20,color=TEAL,bold=True,align='center',group=2)
for row in range(3):
    for col in range(4):
        s.box(369+col*31,545+row*23,24,17,fill=TEAL if (row,col)==(1,2) else '38505E',r=3,group=2)
s.text(355,641,151,31,'Gaze keyboard',16,color=MUTED,align='center',group=2)
s.text(103,713,400,34,'Concept illustration',21,color=MUTED,align='center')
s.text(575,361,918,41,'DRIVE MODE',24,color=TEAL,bold=True,group=1)
s.text(575,410,918,80,'EEG intent or eye selection → Pi → local\nsafety controller → wheelchair motors',29,group=1)
s.text(575,519,918,41,'COMMUNICATION MODE / CHAIR STOPPED',24,color=PURPLE,bold=True,group=2)
s.text(575,569,930,81,'Gaze keyboard → text → Pi hub → Braille,\nscreen, and planned avatar / speech',29,group=2)
s.box(575,689,940,74,fill=PANEL,stroke=ORANGE,group=3)
s.text(596,709,900,41,'E-stop • obstacle stop • signal timeout • speed limit',25,color=ORANGE,bold=True,group=3)

# 11
s=Slide('One hub. One accessible chat.','11 / AI + chat experience',notes='[30 seconds] This is the chat-interface design for our own AI/ML-based hub. The intended workflow combines recognition with receiver selection and accessible delivery. A glove or gaze user composes a message; the hub sends it as Braille, on-screen text or, after integration, avatar or speech. We plan contacts, conversation history, quick requests and device status in one interface. The bubbles shown here are an illustrative conversation, not live device output. Our next engineering milestones are training the custom model, integrating the chat interface and measuring accuracy, false activations and end-to-end delay on unseen users.')
s.head('OUR AI/ML HUB / INTERFACE CONCEPT','One hub. One accessible chat.','Recognize intent. Choose a person. Deliver in their preferred format.')
s.box(80,373,869,389,fill=PANEL,r=24)
s.text(110,397,460,35,'ABILITY CHAT',25,color=TEAL,bold=True)
s.chip(655,390,'INTERFACE CONCEPT',PURPLE,261)
s.line(107,447,921,447,'38505E',2)
s.box(112,471,440,88,fill='304958',r=18,group=1)
s.text(133,480,398,25,'GAZE INPUT → BRAILLE CONTACT',15,color=TEAL,bold=True,group=1)
s.text(133,514,398,38,'Water, please.',29,color=WHITE,group=1)
s.box(493,583,420,76,fill=TEAL,r=18,group=2)
s.text(515,602,376,40,'On my way.',29,color=INK,bold=True,group=2)
s.box(111,693,803,44,fill=NAVY,r=12,group=3)
s.text(127,704,771,29,'Quick requests     |     Contacts     |     Message history',21,color=MUTED,group=3)
for i,(title,body,col) in enumerate([('OUR OWN AI/ML MODEL','Custom recognition + user calibration',TEAL),('ACCESSIBLE CONVERSATIONS','Text, tactile and planned avatar / speech',ORANGE),('A CONNECTED WORKSPACE','Contacts, device status and quick requests',PURPLE)]):
    y=385+i*132
    s.text(994,y,524,39,title,23,color=col,bold=True,group=i+1)
    s.text(994,y+48,524,68,body,27,color=WHITE,group=i+1)
s.text(82,778,1430,28,'Next: train the model, integrate chat, and validate recognition + delivery with users.',19,color=MUTED)

# 12
s=Slide('Different abilities. Shared possibilities.','12 / Closing',notes='[15 seconds] Project Ability brings sensing, embedded control and message routing into one modular design. The current code gives us gesture input, a tactile display, button entry and a hub-routing foundation. Our next challenge is to validate the experience and complete the return paths. The goal is simple: different abilities, shared possibilities. Thank you.')
s.text(80,141,1250,37,'PROJECT ABILITY',24,color=TEAL,bold=True)
s.text(73,242,1430,235,'Different abilities.\nShared possibilities.',87,bold=True)
s.text(81,568,1310,87,'Start with a message.\nBuild toward more independence.',36,color=MUTED)
for x,lab,col in [(83,'MODULAR',TEAL),(406,'AFFORDABLE',ORANGE),(729,'CONNECTED',PURPLE)]:s.chip(x,713,lab,col,278,2)
s.text(1238,718,278,40,'Thank you.',31,color=WHITE,bold=True,align='right')

# 13 — detailed core BOM, backup only. BDT cost targets, not supplier quotes.
BOM=[('Hub','Raspberry Pi 5: 2GB / reused target',1,12000),('Hub','USB-C power supply',1,1500),('Hub','Active cooler',1,500),('Hub','microSD storage',1,600),('Hub','Case',1,400),('Glove','ESP32-S3',1,650),('Glove','Flex sensors: economy target',5,250),('Glove','BNO055: sourcing target',1,1900),('Glove','Glove + wiring / resistors',1,400),('Glove','Power allowance',1,300),('Braille','ESP32-S3',1,650),('Braille','MG90S servos',8,200),('Braille','Cams / pins / housing',1,600),('Braille','Tactile buttons',7,10),('Braille','Regulated 5V supply',1,400),('Braille','Cables / connectors',1,180)]
EXTRA_BOM=[('Voice','ESP32-S3',1,650),('Voice','I2S microphone',1,250),('Voice','Amplifier + speaker',1,350),('Voice','Power + wiring',1,250),('Gaze','USB camera: proposed economy option',1,2000),('Gaze','Mount + illumination allowance',1,500),('Gaze','Power + cabling',1,500),('Gaze','Display mounting; screen reused',1,500),('Wheelchair add-on','ESP32-S3 supervision board',1,650),('Wheelchair add-on','Obstacle sensors: allowance',1,1500),('Wheelchair add-on','E-stop + protected interface allowance',1,5000),('Wheelchair add-on','Power conversion + wiring',1,2000),('Wheelchair add-on','Enclosure + mounting',1,1500),('Wheelchair add-on','Integration allowance',1,1350)]
s=Slide('Core build: ৳23,000 target.','Appendix A / BDT budget',light=True,appendix=True,notes='Backup slide. All amounts are Bangladeshi taka cost targets. The glove total of 4,500 follows the owner’s 4,000–5,000 estimate; its line items are allocations, not actual receipts. Braille totals 3,500; Pi hub 15,000. The proposed lower hub target uses a lower-memory or reused Pi 5, not a newly purchased 8GB kit. Economy sourcing and model-inference performance must be checked. Core total is 23,000. Voice adds 1,500 and gaze adds 3,500, yielding 28,000, or 30,800 with 10 percent reserve. Optional wheelchair interface adds 12,000 before EEG, electrodes, powered chair/motors/brakes/battery and testing. The full wheelchair cost is not established. All component rows, including optional modules, are in budget.csv.')
s.head('BACKUP / COMPONENT COST TARGETS IN BDT','Core build: ৳23,000 target.')
for col,rows in enumerate([BOM[:8],BOM[8:]]):
    x=80+col*740
    s.text(x,308,484,35,'COMPONENT',19,color=GRAY,bold=True)
    s.text(x+491,308,60,35,'QTY',19,color=GRAY,bold=True,align='center')
    s.text(x+566,308,113,35,'BDT',19,color=GRAY,bold=True,align='right')
    for i,(mod,name,q,u) in enumerate(rows):
        y=357+i*43
        s.text(x,y,490,35,name,22,color=INK)
        s.text(x+491,y,60,35,str(q),22,color=GRAY,align='center')
        s.text(x+558,y,121,35,f'{q*u:,.0f}',22,color=INK,align='right')
s.box(80,713,1440,67,fill=INK,r=12)
s.text(105,728,1390,39,'GLOVE ৳4,500     +     BRAILLE ৳3,500     +     SHARED HUB ৳15,000',29,color=WHITE,bold=True,align='center')
s.text(83,786,1430,23,'Voice + gaze: ৳5,000 extra. Optional wheelchair control: ৳12,000 excluding EEG + chair. Full breakdown: budget.csv',16,color=GRAY)

# 14
s=Slide('What the evidence supports','Appendix B / Claims',appendix=True,notes='Backup for judge questions. Static source inspection is not physical validation. See PROJECT_ANALYSIS.md for file-specific findings and limitations. Avoid first-ever, perfect accuracy, full sign-language translation, end-to-end encrypted or all routes implemented claims. Orbit Chat and prior EEG/EOG wheelchair papers establish overlap with broad first-ever statements. The compelling contribution is this particular affordable modular integration, to be validated.')
s.head('BACKUP / JUDGE QUESTIONS','What the evidence supports')
rows=[('SOURCE PRESENT','Glove rules, MQTT routing, 4-cell actuation, button input',TEAL),('PARTIAL INTEGRATION','Braille keys display locally; avatar return route remains',ORANGE),('PLANNED','Custom ML, chat UI, gaze, avatar / speech, EEG wheelchair',PURPLE),('MEASURE NEXT','Gesture macro-F1, false triggers, tactile correctness, delay',TEAL),('NOVELTY POSITION','An affordable modular integration; no world-first claim',ORANGE)]
for i,(label,body,col) in enumerate(rows):
    y=340+i*85
    s.text(80,y,371,45,label,24,color=col,bold=True)
    s.text(478,y,1040,59,body,27)

# 15
SOURCES=[
('Rokoko Smartgloves II · ≈ ৳2.45 lakh · professional pair','https://store.rokoko.com/products/smartgloves-ii'),
('HumanWare Brailliant BI 40X · ≈ ৳4.62 lakh · 40 cells','https://store.humanware.com/hus/brailliant-bi-40x-braille-display.html'),
('Bangladesh Bank · rounded presentation rate: USD 1 ≈ ৳123','https://www.bb.org.bd/en/'),
('Orbit Reader 20 · ≈ ৳98,300 · lower-price market reference','https://www.orbitresearch.com/products/blindness-products/braille-devices/orbit-reader-20/'),
('OpenBCI Cyton · ≈ ৳1.54 lakh · EEG board alone','https://shop.openbci.com/products/cyton-biosensing-board-8-channel'),
('Prior EEG/EOG wheelchair research · 2014','https://pmc.ncbi.nlm.nih.gov/articles/PMC4155067/'),
('Parkinson’s Foundation · vision and eye movement changes','https://www.parkinson.org/understanding-parkinsons/non-movement-symptoms/vision')]
s=Slide('Sources & presentation scope','Appendix C / Sources',light=True,appendix=True,notes='Public sources checked September 30, 2026. Click the source titles in the browser, PDF or PowerPoint. The deck is based on static review of this checkout and the user’s requested vision; no hardware was tested and no clinical or accuracy outcomes are claimed. Longer source notes and price assumptions are in SOURCES.md.')
s.head('BACKUP / CLICKABLE REFERENCES','Sources & presentation scope')
for i,(title,url) in enumerate(SOURCES):
    y=309+i*56
    s.text(82,y,1400,43,f'{i+1:02d}   {title}',26,color=INK,link=url)
s.text(82,732,1420,59,'Checked 30 Sep 2026 · Static code review, not a hardware trial.\nConcept illustrations are original diagrams, not photographs of built hardware.',22,color=GRAY)

def render_html():
    sections=[]
    for i,s in enumerate(slides):
        out=[]
        for e in s.elements:
            pos=f'left:{e["x"]}px;top:{e["y"]}px;'
            cls=f'element g{min(e.get("group",0),5)}'
            if e['kind']=='text':
                style=pos+f'width:{e["w"]}px;height:{e["h"]}px;font-size:{e["size"]}px;color:#{e["color"]};font-weight:{700 if e["bold"] else 400};text-align:{e["align"]};'
                content=html.escape(e['text']).replace('\n','<br>')
                if e.get('link'):content=f'<a href="{html.escape(e["link"])}" target="_blank" rel="noopener">{content}</a>'
                out.append(f'<div class="{cls} text" style="{style}">{content}</div>')
            elif e['kind'] in ('box','circle'):
                style=pos+f'width:{e["w"]}px;height:{e["h"]}px;border-radius:{"50%" if e["kind"]=="circle" else str(e["r"])+"px"};background:{"#"+e["fill"] if e["fill"] else "transparent"};'
                if e.get('stroke'):style+=f'border:{e["sw"]}px solid #{e["stroke"]};'
                out.append(f'<div class="{cls}" style="{style}"></div>')
            elif e['kind']=='flow':
                out.append(f'<svg class="element line flow" width="1600" height="900" aria-hidden="true"><circle r="6" fill="#{e["color"]}"><animateMotion dur="2.8s" repeatCount="indefinite" path="M {e["x"]} {e["y"]} L {e["x2"]} {e["y2"]}"/></circle></svg>')
            else:
                # Full-canvas SVG retains exact line positions.
                out.append(f'<svg class="{cls} line" width="1600" height="900"><line x1="{e["x"]}" y1="{e["y"]}" x2="{e["x2"]}" y2="{e["y2"]}" stroke="#{e["color"]}" stroke-width="{e["sw"]}" stroke-linecap="round"/></svg>')
        sections.append(f'<section class="slide {"active" if i==0 else ""}" data-title="{html.escape(s.title)}" data-appendix="{str(s.appendix).lower()}" aria-label="Slide {i+1}: {html.escape(s.title)}" style="background:#{s.bg}">'+''.join(out)+'</section>')
    payload=json.dumps([dict(title=s.title,notes=s.notes) for s in slides]).replace('</','<\\/')
    template='''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Project Ability — Project Show</title><style>
*{box-sizing:border-box}html,body{margin:0;width:100%;height:100%;overflow:hidden;background:#09121b;font-family:Arial,Helvetica,sans-serif}#stage{width:1600px;height:900px;position:absolute;left:50%;top:50%;transform-origin:center center}.slide{position:absolute;inset:0;display:none;overflow:hidden}.slide.active{display:block}.element{position:absolute}.text{line-height:1.16;white-space:normal;overflow:visible;letter-spacing:-.015em}.text a{color:inherit;text-decoration:underline;text-decoration-thickness:1px;text-underline-offset:5px}.line{left:0;top:0;pointer-events:none}.slide.active .g0{animation:appear .38s both}.slide.active .g1{animation:appear .55s .12s both}.slide.active .g2{animation:appear .55s .28s both}.slide.active .g3{animation:appear .55s .44s both}.slide.active .g4{animation:appear .55s .60s both}.slide.active .g5{animation:appear .55s .76s both}@keyframes appear{from{opacity:0;transform:translateY(14px)}to{opacity:1;transform:translateY(0)}}
#controls{position:fixed;bottom:10px;left:50%;transform:translateX(-50%);display:flex;gap:7px;align-items:center;background:#071019ed;border:1px solid #445864;border-radius:14px;padding:7px;color:white;opacity:.16;transition:opacity .2s;z-index:10}#controls:hover,#controls:focus-within{opacity:1}button,select{font:14px Arial;color:#fff;background:#233746;border:1px solid #526775;border-radius:7px;padding:8px;cursor:pointer}button:focus-visible,select:focus-visible,a:focus-visible{outline:3px solid #77e1cb}#notes{display:none;position:fixed;right:20px;top:20px;width:min(560px,90vw);background:#091722f5;color:#fff;padding:26px;border:1px solid #77e1cb;border-radius:14px;font-size:19px;line-height:1.55;z-index:20;max-height:85vh;overflow:auto}#notes.open{display:block}#blackout{display:none;position:fixed;inset:0;background:black;z-index:30}#blackout.open{display:block}#progress{height:3px;position:fixed;bottom:0;left:0;background:#77e1cb;transition:width .3s;z-index:12}.print-mode #stage{position:static;transform:none!important}.print-mode .slide{position:relative;display:block}.print-mode .element{animation:none!important}.print-mode #controls,.print-mode #notes,.print-mode #progress{display:none}
@media(prefers-reduced-motion:reduce){*{animation:none!important;transition:none!important}.flow{display:none}}@media print{@page{size:1600px 900px;margin:0}html,body{width:1600px;height:auto;overflow:visible;background:white}#stage{position:static;transform:none!important;width:1600px;height:auto}.slide{position:relative;display:block;break-after:page;page-break-after:always;width:1600px;height:900px;-webkit-print-color-adjust:exact;print-color-adjust:exact}.element{animation:none!important}#controls,#notes,#blackout,#progress,.flow{display:none!important}}
</style></head><body><main id="stage">__SLIDES__</main><nav id="controls" aria-label="Presentation controls"><button id="prev" aria-label="Previous slide">←</button><select id="jump" aria-label="Go to slide"></select><button id="next" aria-label="Next slide">→</button><button id="full">Fullscreen · F</button><button id="note">Notes · N</button><button id="black">Blackout · B</button><button id="play">Auto · A</button></nav><aside id="notes" aria-label="Speaker notes"></aside><div id="blackout" aria-label="Presentation blacked out"></div><div id="progress"></div><script>
const DATA=__DATA__;let current=0,auto=null;const all=[...document.querySelectorAll('.slide')],stage=document.querySelector('#stage'),jump=document.querySelector('#jump');DATA.forEach((s,i)=>{let o=document.createElement('option');o.value=i;o.textContent=`${i+1}. ${s.title}`;jump.append(o)});
function fit(){stage.style.transform=`translate(-50%,-50%) scale(${Math.min(innerWidth/1600,innerHeight/900)})`}function show(i){current=Math.max(0,Math.min(all.length-1,i));all.forEach((s,j)=>{s.classList.toggle('active',j===current);s.setAttribute('aria-hidden',j===current?'false':'true')});jump.value=current;document.querySelector('#notes').textContent=DATA[current].notes;document.querySelector('#progress').style.width=`${(current+1)/all.length*100}%`;history.replaceState(null,'',`#${current+1}`);if(auto&&(current===6||current===11))stopAuto()}function stopAuto(){clearInterval(auto);auto=null;document.querySelector('#play').textContent='Auto · A'}function autoplay(){if(auto)return stopAuto();if(current===6||current>=11)return;auto=setInterval(()=>show(current+1),18000);document.querySelector('#play').textContent='Pause · A'}function blackout(){stopAuto();document.querySelector('#blackout').classList.toggle('open')}function fullscreen(){if(document.fullscreenElement)document.exitFullscreen();else document.documentElement.requestFullscreen?.().catch(()=>{})}
document.querySelector('#prev').onclick=()=>show(current-1);document.querySelector('#next').onclick=()=>show(current+1);jump.onchange=()=>show(+jump.value);document.querySelector('#full').onclick=fullscreen;document.querySelector('#note').onclick=()=>document.querySelector('#notes').classList.toggle('open');document.querySelector('#black').onclick=blackout;document.querySelector('#blackout').onclick=blackout;document.querySelector('#play').onclick=autoplay;
addEventListener('keydown',e=>{if(['SELECT','INPUT','TEXTAREA'].includes(document.activeElement.tagName))return;let k=e.key.toLowerCase();if(['arrowright','pagedown',' '].includes(k)){e.preventDefault();show(current+1)}if(['arrowleft','pageup'].includes(k)){e.preventDefault();show(current-1)}if(k==='home')show(0);if(k==='end')show(11);if(k==='f')fullscreen();if(k==='n')document.querySelector('#notes').classList.toggle('open');if(k==='b')blackout();if(k==='a')autoplay();if(k==='escape'){document.querySelector('#notes').classList.remove('open');document.querySelector('#blackout').classList.remove('open');stopAuto()}});let sx=0;addEventListener('touchstart',e=>sx=e.changedTouches[0].screenX,{passive:true});addEventListener('touchend',e=>{let d=e.changedTouches[0].screenX-sx;if(Math.abs(d)>60)show(current+(d<0?1:-1))},{passive:true});addEventListener('resize',fit);fit();show((parseInt(location.hash.slice(1))||1)-1);window.ability={show,data:DATA};
</script></body></html>'''
    (ROOT/'Ability_Presentation.html').write_text(template.replace('__SLIDES__',''.join(sections)).replace('__DATA__',payload))

def render_pptx():
    prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
    prs.core_properties.title='Project Ability — Different ways to communicate. One connection.'
    prs.core_properties.subject='Project show: current prototype, budget and proposed EEG / gaze wheelchair extension'
    prs.core_properties.author='Project Ability'
    for s in slides:
        sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor.from_string(s.bg)
        for e in s.elements:
            u=lambda n: Inches(n/100)
            if e['kind']=='flow':continue
            if e['kind']=='text':
                sh=sl.shapes.add_textbox(u(e['x']),u(e['y']),u(e['w']),u(e['h']))
                tf=sh.text_frame;tf.clear();tf.word_wrap=True
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                for i,line in enumerate(e['text'].split('\n')):
                    p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
                    from pptx.enum.text import PP_ALIGN
                    p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[e['align']]
                    p.space_after=Pt(0);p.space_before=Pt(0);p.line_spacing=1.12
                    run=p.add_run();run.text=line;run.font.name='Arial';run.font.size=Pt(e['size']*.72);run.font.bold=e['bold'];run.font.color.rgb=RGBColor.from_string(e['color'])
                    if e.get('link'):run.hyperlink.address=e['link']
            elif e['kind'] in ('box','circle'):
                typ=MSO_SHAPE.OVAL if e['kind']=='circle' else MSO_SHAPE.ROUNDED_RECTANGLE
                sh=sl.shapes.add_shape(typ,u(e['x']),u(e['y']),u(e['w']),u(e['h']))
                if typ==MSO_SHAPE.ROUNDED_RECTANGLE:sh.adjustments[0]=min(.5,e['r']/min(e['w'],e['h']))
                if e['fill']:sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(e['fill'])
                else:sh.fill.background()
                if e.get('stroke'):sh.line.color.rgb=RGBColor.from_string(e['stroke']);sh.line.width=Pt(e['sw']*.72)
                else:sh.line.fill.background()
            else:
                sh=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,u(e['x']),u(e['y']),u(e['x2']),u(e['y2']))
                sh.line.color.rgb=RGBColor.from_string(e['color']);sh.line.width=Pt(e['sw']*.72)
        sl.notes_slide.notes_text_frame.text=s.notes
        transition=OxmlElement('p:transition');transition.set('spd','med');transition.append(OxmlElement('p:fade'))
        sl._element.append(transition)
    prs.save(ROOT/'Ability_Presentation.pptx')

if __name__=='__main__':
    render_html();render_pptx()
    (ROOT/'slides.json').write_text(json.dumps([s.__dict__ for s in slides],indent=2))
    (ROOT/'SPEAKER_NOTES.md').write_text('# Project Ability — speaker notes\n\n12 presentation slides + 3 optional appendices. Approximately 5 minutes of narration plus 2 minutes 45 seconds of live hardware.\n\n'+'\n\n'.join(f'## {i+1:02d}. {s.title}\n\n{s.notes}' for i,s in enumerate(slides)))
    with (ROOT/'budget.csv').open('w',newline='') as f:
        writer=csv.writer(f);writer.writerow(['Module','Component','Quantity','Unit_BDT_target','Line_BDT_target','Basis'])
        for m,n,q,u in BOM+EXTRA_BOM:writer.writerow([m,n,q,u,q*u,'Cost allocation / target; glove total uses owner estimate; not an invoice'])
        writer.writerow(['CORE TOTAL','Glove + Braille + hub','','',23000,'Cost target; existing screen reused'])
        writer.writerow(['COMMUNICATION TOTAL','Core + voice + gaze','','',28000,'Includes proposed voice/gaze; excludes wheelchair'])
        writer.writerow(['RESERVE','10 percent of communication target','','',2800,'Planning contingency'])
        writer.writerow(['COMMUNICATION WITH RESERVE','Target budget','','',30800,'Not actual expenditure; excludes labor, tax, shipping'])
        writer.writerow(['OPTIONAL ADD-ON','Wheelchair control interface','','',12000,'Excludes EEG, electrodes, chair, motors, brakes, battery and validation'])
    print(f'Built {len(slides)} slides: HTML, editable PPTX, notes, and budget CSV')
